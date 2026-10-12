"""Excel (.xlsx) ingestion: every non-empty worksheet becomes a table.

Handles merged cells (values are propagated across the merged range), header rows that are not on
row 1 (title rows above the table), blank rows/columns, and mixed-type columns.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

import duckdb
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.core.errors import IngestionError
from app.services.ingestion.naming import quote_ident, to_identifier, unique_identifiers
from app.services.ingestion.types import LoadedTable

_HEADER_SCAN_ROWS = 15
_READ_ONLY_THRESHOLD_BYTES = 25 * 1024 * 1024


def _matrix(sheet: Worksheet, unmerge: bool) -> list[list[Any]]:
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    if unmerge and getattr(sheet, "merged_cells", None):
        for merged in list(sheet.merged_cells.ranges):
            top_left = sheet.cell(row=merged.min_row, column=merged.min_col).value
            for r in range(merged.min_row - 1, min(merged.max_row, len(rows))):
                for c in range(merged.min_col - 1, merged.max_col):
                    while len(rows[r]) <= c:
                        rows[r].append(None)
                    rows[r][c] = top_left
    return rows


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def infer_header_row(rows: list[list[Any]]) -> int:
    """Pick the first row that looks like a header: mostly non-empty text, followed by data."""
    best_index, best_score = 0, -1.0
    width = max((len(r) for r in rows[:_HEADER_SCAN_ROWS]), default=0)
    if width == 0:
        return 0
    for index, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        filled = [v for v in row if not _is_blank(v)]
        if len(filled) < max(1, int(0.5 * width)):
            continue
        text_ratio = sum(isinstance(v, str) for v in filled) / len(filled)
        unique_ratio = len({str(v) for v in filled}) / len(filled)
        has_data_below = any(
            any(not _is_blank(v) for v in nxt) for nxt in rows[index + 1 : index + 4]
        )
        score = text_ratio * 2 + unique_ratio + (len(filled) / width) + (1 if has_data_below else -2)
        if score > best_score + 1e-9:
            best_index, best_score = index, score
        if text_ratio == 1.0 and unique_ratio == 1.0 and len(filled) == width and has_data_below:
            break
    return best_index


def _normalise_column(values: list[Any]) -> pd.Series:
    non_null = [v for v in values if not _is_blank(v)]
    if non_null and all(isinstance(v, bool) for v in non_null):
        return pd.Series([None if _is_blank(v) else bool(v) for v in values], dtype="boolean")
    if non_null and all(isinstance(v, int | float) and not isinstance(v, bool) for v in non_null):
        return pd.to_numeric(pd.Series([None if _is_blank(v) else v for v in values]), errors="coerce")
    if non_null and all(isinstance(v, dt.datetime | dt.date) for v in non_null):
        return pd.to_datetime(pd.Series([None if _is_blank(v) else v for v in values]), errors="coerce")
    return pd.Series(
        [None if _is_blank(v) else (v.isoformat() if isinstance(v, dt.date) else str(v).strip()) for v in values],
        dtype="object",
    )


def sheet_to_frame(rows: list[list[Any]]) -> tuple[pd.DataFrame, list[str], int, list[str]]:
    warnings: list[str] = []
    while rows and all(_is_blank(v) for v in rows[-1]):
        rows.pop()
    if not rows:
        return pd.DataFrame(), [], 0, warnings
    header_index = infer_header_row(rows)
    if header_index > 0:
        warnings.append(f"Header detected on row {header_index + 1}; {header_index} row(s) above it were skipped")
    header = rows[header_index]
    body = [r for r in rows[header_index + 1 :] if not all(_is_blank(v) for v in r)]
    dropped_blank = len(rows) - header_index - 1 - len(body)
    if dropped_blank:
        warnings.append(f"{dropped_blank} empty row(s) removed")
    width = max([len(header)] + [len(r) for r in body]) if body else len(header)
    header = list(header) + [None] * (width - len(header))
    keep = [
        c for c in range(width)
        if not _is_blank(header[c]) or any(c < len(r) and not _is_blank(r[c]) for r in body)
    ]
    original = [str(header[c]).strip() if not _is_blank(header[c]) else f"column_{c + 1}" for c in keep]
    columns = {
        str(i): _normalise_column([r[c] if c < len(r) else None for r in body]) for i, c in enumerate(keep)
    }
    frame = pd.DataFrame(columns)
    return frame, original, header_index, warnings


def load_excel(
    con: duckdb.DuckDBPyConnection, path: Path, sheets: list[str] | None = None
) -> list[LoadedTable]:
    read_only = path.stat().st_size > _READ_ONLY_THRESHOLD_BYTES
    try:
        workbook = load_workbook(path, read_only=read_only, data_only=True)
    except Exception as exc:
        raise IngestionError(f"Could not open the Excel workbook: {str(exc)[:200]}") from exc

    loaded: list[LoadedTable] = []
    used_names: set[str] = set()
    try:
        for sheet in workbook.worksheets:
            if sheets and sheet.title not in sheets:
                continue
            rows = _matrix(sheet, unmerge=not read_only)
            frame, original, header_index, warnings = sheet_to_frame(rows)
            if frame.empty:
                continue
            if read_only:
                warnings.append("Large workbook opened in streaming mode; merged cells were not expanded")
            names = unique_identifiers(original)
            frame.columns = names
            table = to_identifier(sheet.title, fallback="sheet")
            while table in used_names:
                table = f"{table}_x"
            used_names.add(table)
            con.register("__frame", frame)
            con.execute(f"CREATE TABLE {quote_ident(table)} AS SELECT * FROM __frame")
            con.unregister("__frame")
            loaded.append(
                LoadedTable(
                    name=table,
                    display_name=sheet.title,
                    original_columns=original,
                    columns=names,
                    row_count=len(frame),
                    warnings=warnings,
                    details={"sheet": sheet.title, "header_row": header_index + 1},
                )
            )
    finally:
        workbook.close()
    if not loaded:
        raise IngestionError("The workbook has no worksheets with data")
    return loaded


def list_sheets(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True)
    try:
        return [ws.title for ws in workbook.worksheets]
    finally:
        workbook.close()

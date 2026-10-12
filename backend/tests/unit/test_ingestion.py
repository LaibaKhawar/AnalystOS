from __future__ import annotations

from pathlib import Path

import duckdb
import openpyxl
import pytest

from app.core.errors import IngestionError
from app.services.ingestion.coercion import coerce_text_columns, reorder_columns
from app.services.ingestion.csv_reader import detect_encoding, load_csv, sniff_structure
from app.services.ingestion.excel_reader import infer_header_row, load_excel
from app.services.ingestion.naming import to_identifier, unique_identifiers
from app.services.ingestion.service import build_dataset_files, detect_file_type


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Revenue ($)", "revenue_usd"),
        ("Customer ID", "customer_id"),
        ("2026 Sales", "col_2026_sales"),
        ("Growth %", "growth_pct"),
        ("orderDate", "order_date"),
        ("  ", "column"),
        ("select", "select_"),
        ("Région", "region"),
        ("date", "date"),
    ],
)
def test_identifiers(raw: str, expected: str) -> None:
    assert to_identifier(raw) == expected


def test_unique_identifiers_deduplicate() -> None:
    assert unique_identifiers(["Sales", "sales", "SALES", ""]) == ["sales", "sales_2", "sales_3", "column_4"]


def _write(tmp: Path, name: str, content: bytes) -> Path:
    path = tmp / name
    path.write_bytes(content)
    return path


def test_csv_semicolon_latin1_currency_malformed(tmp_path: Path) -> None:
    text = (
        'Order Date;Region;Revenue ($);Growth %\n2026-01-03;North;"$1,200.50";12%\n2026-01-04;Sóuth;(300);5%\n'
        "broken;row\n2026-01-05;East;$900;7%\n2026-01-06;West;$1,000;8%;extra\n"
    )
    path = _write(tmp_path, "messy.csv", text.encode("latin-1"))
    assert detect_encoding(path) == "cp1252"
    con = duckdb.connect()
    table = load_csv(con, path, "sales", tmp_path)
    assert table.details["delimiter"] == ";"
    assert table.details["rejected_rows"] == 2
    assert table.columns == ["order_date", "region", "revenue_usd", "growth_pct"]
    coercions = {c.column: c for c in coerce_text_columns(con, "sales")}
    reorder_columns(con, "sales", table.columns)
    assert coercions["revenue_usd"].unit == "currency"
    assert coercions["growth_pct"].unit == "percent"
    rows = con.execute("SELECT region, revenue_usd FROM sales ORDER BY order_date").fetchall()
    assert rows == [("North", 1200.5), ("Sóuth", -300.0), ("East", 900.0)]


def test_csv_without_header(tmp_path: Path) -> None:
    path = _write(tmp_path, "nohdr.csv", b"1,2.5,3\n4,5.5,6\n7,8.5,9\n")
    delimiter, has_header, width = sniff_structure(path)
    assert (delimiter, has_header, width) == (",", False, 3)


def test_csv_utf8_bom_and_nulls(tmp_path: Path) -> None:
    path = _write(tmp_path, "bom.csv", "﻿a,b\n1,NA\n2,\n3,7\n".encode())
    con = duckdb.connect()
    table = load_csv(con, path, "t", tmp_path)
    assert table.columns == ["a", "b"]
    assert con.execute("SELECT count(b) FROM t").fetchone() == (1,)


def test_empty_csv_rejected(tmp_path: Path) -> None:
    path = _write(tmp_path, "empty.csv", b"")
    with pytest.raises(IngestionError):
        load_csv(duckdb.connect(), path, "t", tmp_path)


def test_file_type_detection() -> None:
    assert detect_file_type("x.CSV") == "csv"
    assert detect_file_type("book.xlsx") == "xlsx"
    with pytest.raises(IngestionError, match="xlsx"):
        detect_file_type("old.xls")
    with pytest.raises(IngestionError):
        detect_file_type("data.json")


def test_header_inference_skips_title_rows() -> None:
    rows = [["Quarterly report", None, None], [None, None, None], ["Region", "Month", "Revenue"], ["N", "Jul", 1]]
    assert infer_header_row(rows) == 2


def test_excel_multi_sheet_merged_cells(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q3 Summary"
    ws.append(["Quarterly report"])
    ws.append([])
    ws.append(["Region", "Month", "Revenue"])
    ws.append(["North", "Jul", 100])
    ws.append([None, "Aug", 120])
    ws.merge_cells("A4:A5")
    ws.append([])
    ws.append(["South", "Jul", 90.5])
    other = wb.create_sheet("Products")
    other.append(["SKU", "Price"])
    other.append(["A", 1.5])
    wb.create_sheet("Empty")
    path = tmp_path / "book.xlsx"
    wb.save(path)
    con = duckdb.connect()
    tables = load_excel(con, path)
    assert [t.name for t in tables] == ["q3_summary", "products"]
    assert con.execute("SELECT region, revenue FROM q3_summary ORDER BY revenue").fetchall() == [
        ("South", 90.5),
        ("North", 100.0),
        ("North", 120.0),
    ]
    assert any("Header detected on row 3" in w for w in tables[0].warnings)
    only = load_excel(duckdb.connect(), path, sheets=["Products"])
    assert [t.name for t in only] == ["products"]


def test_build_dataset_files_sample_workbook(tmp_path: Path, data_dir: Path) -> None:
    tables, profiles, _report = build_dataset_files(data_dir / "sales_2026.xlsx", "xlsx", "sales", tmp_path / "out")
    names = {t.name for t in tables}
    assert names == {"transactions", "targets", "regions"}
    assert (tmp_path / "out" / "parquet" / "transactions.parquet").exists()
    transactions = next(p for p in profiles if p.name == "transactions")
    import pandas as pd

    expected = int((pd.read_csv(data_dir / "sales.csv")["date"] >= "2026-01-01").sum())
    assert transactions.row_count == expected
    roles = {c.name: c.role for c in transactions.columns}
    assert roles["date"] == "time" and roles["revenue"] == "metric" and roles["region"] == "dimension"
